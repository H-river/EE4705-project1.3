#!/usr/bin/env python
"""Fill the report draft's [token] placeholders from docs/submission/PLACEHOLDERS.md (0 API calls).

    python scripts/fill_report.py [--draft docs/submission/report_draft.docx] [--out docs/submission/report_filled.docx]

For every paragraph (body, tables, headers, footers) each known "[token]" is replaced by the Fill column of the
PLACEHOLDERS.md table. A token split across several Word runs is merged into one run that keeps the first run's
formatting, and the yellow highlight is removed from the inserted value only. Tokens with an empty Fill, or
bracketed highlighted text that is not in the table, are left untouched and printed at the end.
"""

from __future__ import annotations

import argparse
import copy
import pathlib
import re
import sys

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

ROOT = pathlib.Path(__file__).resolve().parents[1]
TOKEN_RE = re.compile(r"\[[^\[\]]{1,80}\]")


def load_fills(md: pathlib.Path) -> dict[str, str]:
    fills, in_table = {}, False
    for line in md.read_text().splitlines():
        if line.startswith("| Token | Section | Fill"):
            in_table = True
            continue
        if in_table and not line.startswith("|"):
            break
        if in_table and line.startswith("| ["):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if cells[2]:
                fills[cells[0]] = cells[2]
    return fills


def paragraphs(doc):
    def walk(container):
        for p in container.paragraphs:
            yield p
        for t in getattr(container, "tables", []):
            for row in t.rows:
                for cell in row.cells:
                    yield from walk(cell)
    yield from walk(doc)
    for s in doc.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            yield from walk(part)


def replace_in_paragraph(p, token: str, value: str) -> int:
    n = 0
    while True:
        runs = p.runs
        text = "".join(r.text for r in runs)
        i = text.find(token)
        if i < 0:
            return n
        j = i + len(token)
        pos, spans = 0, []
        for r in runs:
            spans.append((r, pos, pos + len(r.text)))
            pos += len(r.text)
        hit = [(r, a, b) for r, a, b in spans if a < j and b > i]
        first, fa, _ = hit[0]
        last, la, _ = hit[-1]
        prefix = first.text[: i - fa]
        suffix = last.text[j - la:]
        for r, _, _ in hit[1:]:
            r.text = ""
        # value in its own run: first run's formatting, no highlight
        val = copy.deepcopy(first._r)
        first._r.addnext(val)
        from docx.text.run import Run
        val_run = Run(val, first._parent)
        val_run.text = value
        val_run.font.highlight_color = None
        first.text = prefix
        if suffix:
            tail = copy.deepcopy(first._r)
            val.addnext(tail)
            tail_run = Run(tail, first._parent)
            tail_run.text = suffix
        n += 1


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--draft", type=pathlib.Path, default=ROOT / "docs/submission/report_draft.docx")
    ap.add_argument("--out", type=pathlib.Path, default=ROOT / "docs/submission/report_filled.docx")
    ap.add_argument("--placeholders", type=pathlib.Path, default=ROOT / "docs/submission/PLACEHOLDERS.md")
    args = ap.parse_args(argv)
    if not args.draft.exists():
        print(f"draft not found: {args.draft}", file=sys.stderr)
        return 2
    fills = load_fills(args.placeholders)
    doc = Document(args.draft)
    done: dict[str, int] = {}
    for p in paragraphs(doc):
        text = "".join(r.text for r in p.runs)
        for tok in set(TOKEN_RE.findall(text)):
            if tok in fills:
                done[tok] = done.get(tok, 0) + replace_in_paragraph(p, tok, fills[tok])
    left = {}
    for p in paragraphs(doc):
        found = {r.text.strip() for r in p.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW and r.text.strip()}
        found |= set(TOKEN_RE.findall("".join(r.text for r in p.runs)))
        for tok in found:
            left[tok] = left.get(tok, 0) + 1
    doc.save(args.out)
    print(f"wrote {args.out}: {sum(done.values())} replacements of {len(done)} tokens")
    unused = sorted(set(fills) - set(done))
    if unused:
        print("table tokens not found in the draft:", ", ".join(unused))
    if left:
        print("NOT FILLED (still bracketed or highlighted):")
        for t, c in sorted(left.items()):
            print(f"  {t}  ×{c}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
